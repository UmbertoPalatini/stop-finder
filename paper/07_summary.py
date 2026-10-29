#!/usr/bin/env python
"""Combine 02 (reference species), 05 (Tox/Wye) and 06 (alternative dsx female exon) into one table
of female-specific termination codons, and write it as TSV and XLSX.
Usage: python 07_summary.py results_dir
"""
import sys, csv, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

R = sys.argv[1]
SP = {"Aedes_aegypti": "Aedes aegypti", "Culex_quinquefasciatus": "Culex quinquefasciatus",
      "Anopheles_gambiae": "Anopheles gambiae", "Toxorhynchites_amboinensis": "Toxorhynchites amboinensis",
      "Wyeomyia_smithii": "Wyeomyia smithii"}
ORDER = list(SP)
RNA = lambda c: c.replace("T", "U")

def rows(p):
    return list(csv.DictReader(open(p), delimiter="\t"))

def seglen(txt):
    return "+".join(x.split("(")[1].split("bp")[0] for x in txt.split(";") if "(" in x and ("*" in x or "*" not in txt))

out = []
for r in rows(f"{R}/02_ref_stops/ref_pairs.tsv"):
    out.append(dict(species=r["species"], gene=r["gene"], isoform="female (annotated)", source="genome + edited annotation",
                    female_aa=r["female_prot_len"], male_aa=r["male_prot_len"],
                    female_segment_bp="+".join(x.split("(")[1].split("bp")[0] for x in r["female_specific_segments"].split(";")),
                    stop=r["female_stop_codon"], in_female_seq=r["stop_in_female_specific"], ctx=r["stop_context_-6_+6"],
                    nt4=r["nt4_after_stop"], next_n=r["codons_to_next_inframe_stop"], next_c=r["next_stop_codon"],
                    inframe=r["female_specific_inframe_stops"], support="", where=f"{r['chrom']}:{r['female_stop_genomic']} ({r['strand']})"))
for r in rows(f"{R}/06_dsx_alt/ref_dsx_alt_exon_stops.tsv"):
    out.append(dict(species=r["species"], gene="dsx", isoform="female alternative exon (5b-type; inferred, not annotated)",
                    source="genome, conserved peptide (tblastn)", female_aa="", male_aa="", female_segment_bp="",
                    stop=r["stop_codon"], in_female_seq="True", ctx=r["stop_context_-6_+6"], nt4=r["nt4_after_stop"],
                    next_n=r["codons_to_next_inframe_stop"], next_c=r["next_stop_codon"], inframe="", support="",
                    where=r["stop_genomic"]))
iso_name = {"MOSN": {"F1": "female"}, "fru": {"F1": "female"},
            "dsx": {"F1": "female (canonical exon)", "F2": "female, second exon (Tox/Wye only)", "F3": "female alternative exon (5b-type)"}}
for r in rows(f"{R}/05_toxwye/toxwye_pairs.tsv"):
    segs = [x for x in r["female_specific_segments_in_F"].split(";") if x]
    key = [x for x in segs if x.endswith("*")] or segs
    out.append(dict(species=r["species"], gene=r["gene"], isoform=iso_name[r["gene"]][r["isoform"]], source="de novo Trinity contigs",
                    female_aa=r["female_prot_len"], male_aa=r["male_prot_len"],
                    female_segment_bp=key[0].split("(")[1].split("bp")[0] if key else "",
                    stop=r["female_stop_codon"], in_female_seq=r["stop_in_female_specific"], ctx=r["stop_context_-6_+6"],
                    nt4=r["nt4_after_stop"], next_n=r["codons_to_next_inframe_stop"], next_c=r["next_stop_codon"],
                    inframe=r["female_specific_inframe_stops"] or "none",
                    support=f"female junction in {r['female_junction_samples(F/M)'].split('/')[0]}/3 female and "
                            f"{r['female_junction_samples(F/M)'].split('/')[1]}/3 male samples",
                    where=r["female_contig"].split("|")[1] + " " + r["female_contig"].split("|")[-1] + f" pos {r['female_stop_pos_in_F']}"))

def note(o):
    if o["in_female_seq"] == "False":
        return "No in-frame stop in the female exon; the exon (non-multiple of 3) shifts the frame and translation ends in the next shared exon"
    return ""
gorder = {"MOSN": 0, "dsx": 1, "fru": 2}
out.sort(key=lambda o: (gorder[o["gene"]], ORDER.index(o["species"]), o["isoform"]))

cols = [("Gene", lambda o: o["gene"]), ("Species", lambda o: SP[o["species"]]), ("Isoform", lambda o: o["isoform"]),
        ("Termination codon (DNA)", lambda o: o["stop"]), ("Termination codon (RNA)", lambda o: RNA(o["stop"])),
        ("Stop lies in female-specific sequence", lambda o: "yes" if o["in_female_seq"] == "True" else "no"),
        ("Nucleotide after stop (+4)", lambda o: o["nt4"]), ("Context (-6 to +9)", lambda o: o["ctx"]),
        ("Codons to next in-frame stop", lambda o: o["next_n"]), ("Next in-frame stop (RNA)", lambda o: RNA(o["next_c"])),
        ("In-frame stops in female-specific sequence", lambda o: o["inframe"]),
        ("Female protein (aa)", lambda o: o["female_aa"]), ("Male protein (aa)", lambda o: o["male_aa"]),
        ("Female-specific sequence (bp)", lambda o: o["female_segment_bp"]),
        ("Sex support (Tox/Wye)", lambda o: o["support"]), ("Stop position", lambda o: o["where"]),
        ("Source", lambda o: o["source"]), ("Note", note)]
with open(f"{R}/stop_codon_summary.tsv", "w") as fh:
    fh.write("\t".join(c for c, _ in cols) + "\n")
    for o in out:
        fh.write("\t".join(str(f(o)) for _, f in cols) + "\n")

wb = Workbook()
ws = wb.active; ws.title = "Summary"
ws.append([c for c, _ in cols])
for o in out:
    ws.append([f(o) for _, f in cols])
for c in ws[1]:
    c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDDDDD"); c.alignment = Alignment(wrap_text=True, vertical="top")
for i, w in enumerate([7, 26, 34, 10, 10, 12, 9, 18, 10, 10, 22, 9, 9, 12, 36, 34, 26, 60], 1):
    ws.column_dimensions[ws.cell(1, i).column_letter].width = w
ws.freeze_panes = "D2"
for name, path in (("Ref_all_stops_in_F_seq", "02_ref_stops/ref_female_specific_stops.tsv"),
                   ("ToxWye_all_stops_in_F_seq", "05_toxwye/toxwye_female_specific_stops.tsv"),
                   ("ToxWye_junction_support", "05_toxwye/toxwye_junction_support.tsv"),
                   ("ToxWye_pairs_detail", "05_toxwye/toxwye_pairs.tsv"),
                   ("Ref_pairs_detail", "02_ref_stops/ref_pairs.tsv")):
    w = wb.create_sheet(name)
    for line in open(f"{R}/{path}"):
        w.append(line.rstrip("\n").split("\t"))
    for c in w[1]:
        c.font = Font(bold=True)
wb.save(f"{R}/stop_codon_summary.xlsx")
print(open(f"{R}/stop_codon_summary.tsv").read())
