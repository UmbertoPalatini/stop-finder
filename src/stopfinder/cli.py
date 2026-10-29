"""Command-line interface: stopfinder {transcripts, genome, junctions}."""
import argparse
import csv
import os
import sys

from . import __version__, genome, junctions, transcripts
from .annotation import load_transcripts
from .core import JUNCTION_COLUMNS, STOP_COLUMNS, SUMMARY_COLUMNS, analyse
from .seqio import IndexedFasta, read_fasta


def read_table(path, required):
    with open(path) as fh:
        rows = list(csv.DictReader((l for l in fh if l.strip() and not l.startswith("#")), delimiter="\t"))
    if rows:
        missing = [c for c in required if c not in rows[0]]
        if missing:
            sys.exit(f"{path}: missing column(s) {', '.join(missing)}; found {', '.join(rows[0])}")
    for i, r in enumerate(rows, 2):
        if None in r:
            sys.exit(f"{path}: a row has more columns than the header (row {i}, after skipping comments)")
        for k, v in list(r.items()):
            r[k] = (v or "").strip()
    return rows


def opt_int(v):
    return int(v.replace(",", "")) if v not in (None, "", ".", "NA") else None


def write_tsv(path, cols, rows):
    with open(path, "w", newline="") as fh:
        w = csv.DictWriter(fh, cols, delimiter="\t", extrasaction="ignore", lineterminator="\n")
        w.writeheader()
        w.writerows(rows)


def write_xlsx(path, sheets):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        sys.exit("--xlsx needs openpyxl (pip install openpyxl)")
    wb = Workbook()
    wb.remove(wb.active)
    for name, cols, rows in sheets:
        ws = wb.create_sheet(name)
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([r.get(c, "") for c in cols])
        ws.freeze_panes = "B2"
    wb.save(path)


def run_pairs(pairs, outdir, table, xlsx):
    """pairs: iterable of (pair_id, callable returning PairInput). Runs the analysis and writes outputs."""
    os.makedirs(outdir, exist_ok=True)
    summary, stops, juncs, failed = [], [], [], 0
    with open(os.path.join(outdir, "proteins.fa"), "w") as fp, open(os.path.join(outdir, "transcripts.fa"), "w") as ft:
        for pid, build in pairs:
            try:
                p = build()
                row, per_stop, jn, (fprot, mprot) = analyse(p, table)
            except (ValueError, KeyError) as e:
                failed += 1
                msg = str(e).strip("'\"")
                print(f"stopfinder: {pid}: {msg}", file=sys.stderr)
                summary.append({"pair_id": pid, "note": f"ERROR: {msg}"})
                continue
            summary.append(row)
            stops += per_stop
            juncs += jn
            fp.write(f">{pid}|female|{p.female_id}\n{fprot}\n>{pid}|male|{p.male_id}\n{mprot}\n")
            ft.write(f">{pid}|female|{p.female_id} CDS={row['female_cds']} "
                     f"female_specific={row['female_specific_segments'].replace(' ', '_') or 'none'}\n{p.female}\n")
            ft.write(f">{pid}|male|{p.male_id} CDS={row['male_cds']}\n{p.male}\n")
    write_tsv(os.path.join(outdir, "summary.tsv"), SUMMARY_COLUMNS, summary)
    write_tsv(os.path.join(outdir, "female_specific_stops.tsv"), STOP_COLUMNS, stops)
    write_tsv(os.path.join(outdir, "junctions.tsv"), JUNCTION_COLUMNS, juncs)
    if xlsx:
        write_xlsx(os.path.join(outdir, "summary.xlsx"),
                   [("summary", SUMMARY_COLUMNS, summary), ("female_specific_stops", STOP_COLUMNS, stops)])
    ok = len(summary) - failed
    print(f"stopfinder: {ok} pair(s) analysed, {failed} failed; results in {outdir}/", file=sys.stderr)
    return 1 if failed else 0


def cmd_transcripts(a):
    if a.pairs:
        if a.female or a.male:
            sys.exit("use either --female/--male or --pairs/--fasta, not both")
        if not a.fasta:
            sys.exit("--pairs needs --fasta")
        seqs, headers = {}, {}
        for path in a.fasta:
            s, h = read_fasta(path)
            for k in s:
                if k in seqs and seqs[k] != s[k]:
                    sys.exit(f"sequence id '{k}' appears in more than one FASTA with different sequences; "
                             "make the IDs unique (for example prefix them with the sample name)")
            seqs.update(s)
            headers.update(h)
        rows = read_table(a.pairs, ["pair_id", "female", "male"])
        fseqs = mseqs = seqs
        fheads = mheads = headers
    elif a.female and a.male:
        # read separately: a female and a male contig can share a name (e.g. TRINITY_DN1_c0_g1_i1)
        fseqs, fheads = read_fasta(a.female)
        mseqs, mheads = read_fasta(a.male)
        if len(fseqs) != 1 or len(mseqs) != 1:
            sys.exit("--female and --male must contain one sequence each; use --pairs for several pairs")
        rows = [{"pair_id": a.name, "female": next(iter(fseqs)), "male": next(iter(mseqs)),
                 "male_start": str(a.male_start or ""), "female_start": str(a.female_start or "")}]
    else:
        sys.exit("give --female and --male, or --pairs with --fasta")

    def builder(r):
        def build():
            for col, d in (("female", fseqs), ("male", mseqs)):
                if r[col] not in d:
                    raise ValueError(f"sequence '{r[col]}' not found in the FASTA input")
            return transcripts.build_pair(
                r["pair_id"], r["female"], fseqs[r["female"]], r["male"], mseqs[r["male"]],
                male_start=opt_int(r.get("male_start")), female_start=opt_int(r.get("female_start")),
                min_segment=a.min_segment, table=a.table,
                male_header=mheads.get(r["male"], ""), female_header=fheads.get(r["female"], ""))
        return build

    return run_pairs(((r["pair_id"], builder(r)) for r in rows), a.outdir, a.table, a.xlsx)


def cmd_genome(a):
    rows = read_table(a.pairs, ["pair_id", "female_tx", "male_tx"])
    groups = {}
    for r in rows:
        g, ann = r.get("genome") or a.genome, r.get("annotation") or a.annotation
        if not g or not ann:
            sys.exit(f"{r['pair_id']}: no genome/annotation (give --genome/--annotation or table columns)")
        groups.setdefault((g, ann), []).append(r)
    pairs = []
    for (g, ann), rs in groups.items():
        fasta = IndexedFasta(g)
        tx = load_transcripts(ann, {r[c] for r in rs for c in ("female_tx", "male_tx")})
        for r in rs:
            def build(r=r, fasta=fasta, tx=tx):
                for c in ("female_tx", "male_tx"):
                    if r[c] not in tx:
                        raise ValueError(f"transcript '{r[c]}' not found in {ann}")
                return genome.build_pair(r["pair_id"], fasta, tx[r["female_tx"]], tx[r["male_tx"]],
                                         extra_exons=r.get("extra_female_exons", ""), start=opt_int(r.get("start")))
            pairs.append((r["pair_id"], build))
    return run_pairs(pairs, a.outdir, a.table, a.xlsx)


def cmd_junctions(a):
    probes = read_table(a.junctions, JUNCTION_COLUMNS)
    samples = read_table(a.samples, ["sample", "group", "fasta"])
    base = os.path.dirname(os.path.abspath(a.samples))
    for s in samples:
        if not os.path.isabs(s["fasta"]):
            s["fasta"] = os.path.join(base, s["fasta"])
    rows = junctions.count(probes, samples, a.max_mismatches)
    os.makedirs(a.outdir, exist_ok=True)
    write_tsv(os.path.join(a.outdir, "junction_support.tsv"),
              ["pair_id", "junction", "sample", "group", "n_sequences"], rows)
    summ = junctions.summarise(rows)
    write_tsv(os.path.join(a.outdir, "junction_support_summary.tsv"),
              ["pair_id", "junction", "samples_with_junction"], summ)
    for r in summ:
        print(f"{r['pair_id']}\t{r['junction']}\t{r['samples_with_junction']}")
    return 0


def main(argv=None):
    ap = argparse.ArgumentParser(prog="stopfinder", description=(
        "Termination codons in female-specific (sex-specifically spliced) sequence. Compares a female and a "
        "male isoform, translates the female isoform from the shared start codon to its first stop, and "
        "reports where that stop lies, its context and the next in-frame stop."))
    ap.add_argument("--version", action="version", version=f"%(prog)s {__version__}")
    sub = ap.add_subparsers(dest="cmd", required=True)

    def common(p):
        p.add_argument("-o", "--outdir", required=True, help="output folder")
        p.add_argument("--table", type=int, default=1, help="NCBI genetic code (default 1, standard)")
        p.add_argument("--xlsx", action="store_true", help="also write summary.xlsx (needs openpyxl)")

    t = sub.add_parser("transcripts", help="female and male mRNA sequences (FASTA), from any source")
    t.add_argument("--female", help="FASTA with one female mRNA")
    t.add_argument("--male", help="FASTA with one male mRNA")
    t.add_argument("--name", default="pair1", help="pair name when using --female/--male (default pair1)")
    t.add_argument("--male-start", type=int, help="1-based position of the male start codon")
    t.add_argument("--female-start", type=int, help="1-based position of the female start codon")
    t.add_argument("--pairs", help="TSV: pair_id, female, male [, male_start, female_start]")
    t.add_argument("--fasta", action="append", help="FASTA with the sequences named in --pairs (repeatable)")
    t.add_argument("--min-segment", type=int, default=15,
                   help="minimum length (nt) of an unaligned female region to call it female-specific (default 15)")
    common(t)
    t.set_defaults(func=cmd_transcripts)

    g = sub.add_parser("genome", help="female and male transcript IDs in a genome annotation (GTF/GFF3)")
    g.add_argument("--genome", help="genome FASTA (uncompressed; a .fai index is created if missing)")
    g.add_argument("--annotation", help="GTF or GFF3 with the female and male transcripts")
    g.add_argument("--pairs", required=True, help="TSV: pair_id, female_tx, male_tx "
                   "[, extra_female_exons, start, genome, annotation]")
    common(g)
    g.set_defaults(func=cmd_genome)

    j = sub.add_parser("junctions", help="count assembled sequences per sample that contain each junction")
    j.add_argument("--junctions", required=True, help="junctions.tsv written by 'transcripts' or 'genome'")
    j.add_argument("--samples", required=True, help="TSV: sample, group, fasta")
    j.add_argument("--max-mismatches", type=int, default=3, help="mismatches allowed in a 60-nt probe (default 3)")
    j.add_argument("-o", "--outdir", required=True)
    j.set_defaults(func=cmd_junctions)

    a = ap.parse_args(argv)
    return a.func(a)


if __name__ == "__main__":
    sys.exit(main())
